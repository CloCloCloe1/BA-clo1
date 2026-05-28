from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"

INK = "1D1D1F"
MUTED = "6E6E73"
BLUE = "0071E3"
LIGHT_BLUE = "EAF3FF"
LIGHT_GRAY = "F5F5F7"
BORDER = "DADCE0"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_run(paragraph, text, bold=False, color=INK, size=11):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.size = Pt(size)
    return run


def add_title_page(doc: Document, title: str, subtitle: str, audience: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(54)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "BA CONSIGNMENT SYSTEM", bold=True, color=BLUE, size=10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    add_run(p, title, bold=True, color=INK, size=30)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    add_run(p, subtitle, color=MUTED, size=12)

    table = doc.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    set_table_width(table, [1.55, 4.45])
    rows = [
        ("Audience", audience),
        ("App Link", "[APP LINK]"),
        ("Version", "MVP user guide"),
    ]
    for r, (label, value) in enumerate(rows):
        table.cell(r, 0).text = label
        table.cell(r, 1).text = value
        set_cell_shading(table.cell(r, 0), LIGHT_BLUE)
        for cell in table.rows[r].cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(10.5)
                    run.font.color.rgb = RGBColor.from_string(INK)
                    if cell is table.cell(r, 0):
                        run.bold = True

    doc.add_section(WD_SECTION.NEW_PAGE)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.add_run(text)


def add_bullets(doc: Document, items) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_steps_table(doc: Document, steps) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_width(table, [0.7, 5.8])
    table.cell(0, 0).text = "Step"
    table.cell(0, 1).text = "Action"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, BLUE)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.bold = True
                run.font.size = Pt(10)
    for idx, action in enumerate(steps, start=1):
        row = table.add_row()
        row.cells[0].text = str(idx)
        row.cells[1].text = action
        set_cell_shading(row.cells[0], LIGHT_GRAY)
        for cell in row.cells:
            set_cell_margins(cell)
            set_cell_border(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)
                    run.font.color.rgb = RGBColor.from_string(INK)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_two_col_table(doc: Document, rows, widths=(2.0, 4.5), header=None) -> None:
    table = doc.add_table(rows=1 if header else 0, cols=2)
    table.style = "Table Grid"
    set_table_width(table, list(widths))
    if header:
        table.cell(0, 0).text = header[0]
        table.cell(0, 1).text = header[1]
        for cell in table.rows[0].cells:
            set_cell_shading(cell, BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
    for label, value in rows:
        row = table.add_row()
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for cell in row.cells:
            set_cell_margins(cell)
            set_cell_border(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)
                    run.font.color.rgb = RGBColor.from_string(INK)
        for run in row.cells[0].paragraphs[0].runs:
            run.bold = True


def add_note(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    cell.text = text
    set_cell_shading(cell, "FFF8E6")
    set_cell_border(cell, "E4C46A")
    set_cell_margins(cell, 150, 180, 150, 180)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor.from_string(INK)


def setup_styles(doc: Document) -> None:
    section = doc.sections[0]
    for sec in doc.sections:
        sec.top_margin = Inches(0.82)
        sec.bottom_margin = Inches(0.82)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ["List Bullet", "List Number"]:
        style = styles[list_style]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)


def footer(doc: Document, label: str) -> None:
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.text = label
        for run in p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(MUTED)


def build_ba_doc() -> None:
    doc = Document()
    setup_styles(doc)
    add_title_page(
        doc,
        "BA User SOP",
        "How to register, search products, submit records, and edit today's entries.",
        "BA Users",
    )

    add_heading(doc, "Quick Overview")
    add_body(
        doc,
        "Use the BA Consignment System to record tester requests, damages, theft, and restock requests. "
        "Product details are filled automatically after you scan a full barcode or enter the last 6 digits.",
    )
    add_two_col_table(
        doc,
        [
            ("App link", "[APP LINK]"),
            ("Store", "Customer or channel, such as Miniso or TNT"),
            ("Location", "Specific store/location, such as BRO or STL"),
            ("Editable today", "Your own same-day records only"),
        ],
    )

    add_heading(doc, "Register A BA Account")
    add_steps_table(
        doc,
        [
            "Open the app link.",
            "Choose Register BA.",
            "Enter the company registration code.",
            "Create a username and password.",
            "Enter your BA name.",
            "Enter Store, such as Miniso or TNT.",
            "Enter Location, such as BRO or STL.",
            "Click Create BA account, then return to Sign in.",
        ],
    )

    add_heading(doc, "Submit A Record")
    add_steps_table(
        doc,
        [
            "Go to Record.",
            "Select the type: Tester, Damage, Theft, or Restock.",
            "Confirm date, quantity, BA name, Store, and Location.",
            "Scan the full barcode or enter the last 6 digits.",
            "Confirm the product match. If multiple products appear, select the correct one.",
            "Add notes if needed.",
            "Click Save record. The button briefly changes state to prevent duplicate clicks.",
        ],
    )

    add_heading(doc, "Edit Today's Entries")
    add_body(doc, "The Today's entries section shows records submitted today by your own login.")
    add_two_col_table(
        doc,
        [
            ("Editable fields", "Date, Type, BA Name, Qty, Store, Location, Notes"),
            ("Locked fields", "Last6, Full Barcode, Product Name, Brand, Category"),
            ("Previous-day mistakes", "Contact Admin. BA users can only edit same-day records."),
        ],
        header=("Area", "Details"),
    )
    add_note(
        doc,
        "Important: Product fields are locked to protect the product master data. If the wrong product was selected, submit a corrected record and notify Admin.",
    )

    add_heading(doc, "Search Products")
    add_bullets(
        doc,
        [
            "Search by product keyword, full barcode, last 6 digits, brand, or category.",
            "Use the category dropdown to narrow results.",
            "Always confirm product name and brand before saving a record.",
        ],
    )

    add_heading(doc, "Good Submission Habits")
    add_bullets(
        doc,
        [
            "Check Store and Location before saving.",
            "Do not click Save repeatedly.",
            "Use Notes for anything Admin should know.",
            "Review Today's entries before ending your shift.",
        ],
    )

    footer(doc, "BA Consignment System | BA User SOP")
    doc.save(OUT / "BA_User_SOP.docx")


def build_admin_doc() -> None:
    doc = Document()
    setup_styles(doc)
    add_title_page(
        doc,
        "Admin User SOP",
        "How to review submissions, export reports, manage BA accounts, and maintain access settings.",
        "Admin Users",
    )

    add_heading(doc, "Quick Overview")
    add_body(
        doc,
        "Admin users review all BA submissions, filter records by location, and export Excel reports. "
        "Admin users do not use the BA record entry page.",
    )
    add_two_col_table(
        doc,
        [
            ("App link", "[APP LINK]"),
            ("Admin username", "admin"),
            ("Password location", "Managed in Streamlit Secrets"),
            ("Record storage", "Supabase table: records"),
            ("BA account storage", "Supabase table: ba_users"),
        ],
    )

    add_heading(doc, "Review Records")
    add_steps_table(
        doc,
        [
            "Sign in with the admin account.",
            "Open the Admin page.",
            "Review the dashboard metrics: Total, Tester, Damage, Theft, and Restock.",
            "Use the Location dropdown to review a specific location, such as BRO or STL.",
            "Review the table before exporting.",
        ],
    )

    add_heading(doc, "Export Excel")
    add_steps_table(
        doc,
        [
            "Go to Admin.",
            "Select All locations or a specific Location.",
            "Click Export selected location Excel.",
            "Open the downloaded workbook and review the separate report sheets.",
        ],
    )
    add_two_col_table(
        doc,
        [
            ("Sheets included", "Tester, Damage, Theft, Restock, All Records"),
            ("Location exports", "Use Download each location separately when separate files are needed."),
            ("File review", "Check date, BA name, product, quantity, Store, Location, and Notes before sharing."),
        ],
        header=("Export item", "Details"),
    )

    add_heading(doc, "Excel Column Order")
    add_bullets(
        doc,
        [
            "Date",
            "BA Name",
            "Last6",
            "Full Barcode",
            "Product Name",
            "Brand",
            "Qty",
            "Category",
            "Store Name",
            "Location",
            "Report Type",
            "Notes",
        ],
    )

    add_heading(doc, "Manage BA Accounts")
    add_body(doc, "BA users registered through the app are stored in Supabase table ba_users.")
    add_two_col_table(
        doc,
        [
            ("View accounts", "Supabase > Table Editor > ba_users"),
            ("Deactivate account", "Set active = false for the username."),
            ("Delete test account", "Delete the row from ba_users."),
            ("Registration code", "Managed in Streamlit Secrets as REGISTRATION_CODE."),
        ],
        header=("Task", "Where / how"),
    )

    add_heading(doc, "Useful SQL")
    add_body(doc, "Deactivate a BA account:")
    add_body(doc, "update public.ba_users set active = false where username = 'ba_username';")
    add_body(doc, "Delete a test BA account:")
    add_body(doc, "delete from public.ba_users where username = 'ba_username';")

    add_heading(doc, "Common Admin Checks")
    add_bullets(
        doc,
        [
            "Confirm new records appear in the records table.",
            "Confirm Store and Location values are entered consistently.",
            "Keep the registration code internal.",
            "Ask BAs to correct same-day mistakes in Today's entries.",
            "Correct previous-day issues in Supabase or request a corrected submission.",
        ],
    )

    footer(doc, "BA Consignment System | Admin User SOP")
    doc.save(OUT / "Admin_User_SOP.docx")


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_ba_doc()
    build_admin_doc()
    print(OUT / "BA_User_SOP.docx")
    print(OUT / "Admin_User_SOP.docx")
